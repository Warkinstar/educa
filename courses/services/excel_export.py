from io import BytesIO
from django.utils import timezone
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side


def generate_course_analytics_excel(analytics_data):
  course = analytics_data['course']
  students_summary = analytics_data['students_summary']

  wb = openpyxl.Workbook()
  ws = wb.active
  ws.title = 'Успеваемость'

  title_font = Font(name='Calibri', size=15, bold=True, color='1F497D')
  meta_font = Font(name='Calibri', size=10, italic=True, color='595959')
  header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFF')
  header_fill = PatternFill(
      start_color='2F5597', end_color='2F5597', fill_type='solid'
  )

  thin_border = Border(
      left=Side(style='thin', color='D9D9D9'),
      right=Side(style='thin', color='D9D9D9'),
      top=Side(style='thin', color='D9D9D9'),
      bottom=Side(style='thin', color='D9D9D9'),
  )
  center_align = Alignment(horizontal='center', vertical='center')
  left_align = Alignment(horizontal='left', vertical='center')

  # Заголовок курса (объединяем по ширине колонок A-J)
  ws.merge_cells('A1:J1')
  ws['A1'] = f'Отчёт по успеваемости курса: {course.title}'
  ws['A1'].font = title_font
  ws['A1'].alignment = left_align

  # Метаданные (объединяем по ширине колонок A-J)
  ws.merge_cells('A2:J2')
  ws['A2'] = (
      f'Преподаватель: {course.owner.get_full_name() or course.owner.username}'
      f' | Дата выгрузки: {timezone.now().strftime("%d.%m.%Y %H:%M")} |'
      f' Студентов: {analytics_data["total_students_count"]} | Средний'
      f' прогресс: {analytics_data["group_avg_progress"]}%'
  )
  ws['A2'].font = meta_font
  ws['A2'].alignment = left_align

  headers = [
      '№',
      'ФИО студента',
      'Логин',
      'Email',
      'Прогресс курса (%)',
      'Статус курса',
      'Сдано тестов',
      'Средний балл',
      'Всего попыток',
      'Последняя активность',
  ]

  header_row = 4
  for col_idx, header_text in enumerate(headers, start=1):
    cell = ws.cell(row=header_row, column=col_idx, value=header_text)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = center_align

  current_row = 5
  for idx, student in enumerate(students_summary, start=1):
    status_text = 'Завершён' if student['is_completed'] else 'В процессе'
    quizzes_text = (
        f"{student['passed_quizzes_count']} из {student['total_quizzes_count']}"
    )
    avg_score_text = (
        student['avg_score'] if student['avg_score'] is not None else '—'
    )
    last_activity_text = (
        student['last_activity'].strftime('%d.%m.%Y %H:%M')
        if student['last_activity']
        else '—'
    )

    row_values = [
        idx,
        student['full_name'],
        student['username'],
        student['email'],
        f"{student['progress_percent']}%",
        status_text,
        quizzes_text,
        avg_score_text,
        student['total_attempts'],
        last_activity_text,
    ]

    for col_idx, value in enumerate(row_values, start=1):
      cell = ws.cell(row=current_row, column=col_idx, value=value)
      cell.border = thin_border
      cell.alignment = (
          center_align if col_idx in [1, 5, 6, 7, 8, 9, 10] else left_align
      )

    current_row += 1

  # Автоподбор ширины: учитываем только строки таблицы (строка 4 и ниже)
  for col in ws.columns:
    col_letter = openpyxl.utils.get_column_letter(col[0].column)
    # Пропускаем первые 3 строки шапки при расчёте ширины
    table_cells = [cell for cell in col if cell.row >= 4]
    if table_cells:
      max_len = max(len(str(c.value or '')) for c in table_cells)
      ws.column_dimensions[col_letter].width = max(max_len + 4, 8)

  output = BytesIO()
  wb.save(output)
  output.seek(0)
  return output